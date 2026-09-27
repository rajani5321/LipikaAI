import os
import re
from pathlib import Path
from typing import Dict, Any, Tuple
import pypdf
import docx

class DocumentParserService:
    """Service to parse and extract text from various document formats (PDF, DOCX, TXT)."""

    @staticmethod
    def clean_text(text: str) -> str:
        """Sanitizes extracted document text for downstream NLP and ML analysis."""
        if not text:
            return ""
        
        # Replace odd unicode whitespaces, zero-width spaces, and control characters
        text = text.replace('\u200b', '').replace('\ufeff', '')
        text = re.sub(r'[\r\t]', ' ', text)
        
        # Standardize bullet characters
        text = re.sub(r'[\u2022\u2023\u25E6\u2043\u2219▪•◆●]', '\n• ', text)
        
        # Collapse multiple spaces into single space, but preserve newlines
        lines = [re.sub(r'[ ]+', ' ', line).strip() for line in text.split('\n')]
        
        # Remove consecutive empty lines (leave maximum 1 empty line between blocks)
        cleaned_lines = []
        last_empty = False
        for line in lines:
            if not line:
                if not last_empty:
                    cleaned_lines.append("")
                    last_empty = True
            else:
                cleaned_lines.append(line)
                last_empty = False
                
        return "\n".join(cleaned_lines).strip()

    def parse_pdf(self, file_path: str) -> Tuple[str, int]:
        """Extracts text and page count from a PDF file using pypdf."""
        extracted_text = []
        page_count = 0
        
        try:
            reader = pypdf.PdfReader(file_path)
            page_count = len(reader.pages)
            for page_idx, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    extracted_text.append(page_text)
        except Exception as e:
            raise ValueError(f"Failed to parse PDF document: {str(e)}")

        raw_text = "\n".join(extracted_text)
        return self.clean_text(raw_text), page_count

    def parse_docx(self, file_path: str) -> Tuple[str, int]:
        """Extracts text from a DOCX file using python-docx (paragraphs and tables)."""
        extracted_text = []
        try:
            doc = docx.Document(file_path)
            
            # Paragraphs
            for p in doc.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    extracted_text.append(p_text)
            
            # Tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        extracted_text.append(" | ".join(row_text))
                        
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX document: {str(e)}")

        raw_text = "\n".join(extracted_text)
        return self.clean_text(raw_text), 1

    def parse_txt(self, file_path: str) -> Tuple[str, int]:
        """Extracts text from a plain TXT file."""
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
            return self.clean_text(text), 1
        except Exception as e:
            raise ValueError(f"Failed to read TXT file: {str(e)}")

    def extract_document(self, file_path: str) -> Dict[str, Any]:
        """Unified method to detect extension and extract cleaned text and metadata."""
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        ext = path.suffix.lower()
        if ext == ".pdf":
            text, pages = self.parse_pdf(file_path)
        elif ext == ".docx":
            text, pages = self.parse_docx(file_path)
        elif ext == ".txt":
            text, pages = self.parse_txt(file_path)
        else:
            raise ValueError(f"Unsupported file format: {ext}. Only PDF, DOCX, and TXT are supported.")

        word_count = len(text.split())
        return {
            "text": text,
            "page_count": pages,
            "word_count": word_count,
            "filename": path.name,
            "file_size_kb": round(path.stat().st_size / 1024, 2)
        }

document_parser = DocumentParserService()
